"""The New York Fed's natural-rate estimates: Holston-Laubach-Williams r*, as each quarter published it.

`Hlw` reads the real-time vintages file linked from
newyorkfed.org/research/policy/rstar, ``Holston_Laubach_Williams_real_time_estimates.xlsx``.
It has one sheet per vintage, named by its last quarter ("2015Q4"), each the
whole one-sided history as estimated on the data of that quarter. 2015Q4-2020Q2
are the HLW (2017) model; 2022Q4 on are HLW (2023). There are no vintages
2020Q3-2022Q3: the NY Fed suspended the estimates over COVID-19 and resumed with
2022Q4. A reader that holds a vintage until the next keeps 2020Q2 through the gap.

One observation per vintage: its last quarter's r* (the column under the
"Natural Rate (r*)" group whose country header is the series' economy, `SERIES`),
dated that quarter's first day. It is the number the NY Fed published that
quarter. Each sheet's earlier rows are that vintage's estimates of earlier
quarters, on data revised since: not what was known then, and not used. The two
models lay their groups out in a different order (the 2017 model's output gap
comes first, the 2023 model's last), so the column is found by its headers,
never by position. Each sheet's "Final data point" line is not trusted (2019Q2's
says 2019Q1): the last dated row must be the sheet's own quarter, or the parse
fails.

``published`` = the quarter's last day plus `lag_days` calendar days, or the
release day where that came later. HLW keeps its own schedule, tied to Canada's
GDP release; the LW schedule above it on the page follows US GDP and runs up to
four days earlier. The HLW schedules on the page's archived copies (2018 from
August, 2019, 2020, 2025, 2026) put a release 58-65 days after the quarter, the
latest 2019Q2 on 2019-09-03, so 65 is never early for them. Three were later
(`RELEASED_LATE`). The first HLW (2023) vintage came out when the estimates
resumed, on 2023-05-19. The 2025 federal shutdown delayed GDP, so 2025Q3 came
out 2025-12-29 (90 days) and 2025Q4 2026-03-20 (79 days). The file does not
date its releases, and the schedules for 2016 to mid-2018 and 2023-24 are not on
the archived copies, so a release more than 65 days out in those years would go
unseen.
"""

import datetime
import io
import re
import openpyxl
import pandas as pd
import requests
from policypath.sources.base import Source

DATA = "https://www.newyorkfed.org/medialibrary/media/research/economists/williams/data/"
HEADERS = {"User-Agent": "Mozilla/5.0 (policypath research)"}
# Per series: the workbook, the group of columns, and the economy under it.
SERIES = {"HLW_RSTAR": ("Holston_Laubach_Williams_real_time_estimates.xlsx", "Natural Rate (r*)", "US")}
# Vintages released later than their quarter's end plus the usual lag: vintage -> release day.
RELEASED_LATE = {"2022Q4": "2023-05-19", "2025Q3": "2025-12-29", "2025Q4": "2026-03-20"}


def _get(url):
    r = requests.get(url, headers=HEADERS, timeout=120)
    r.raise_for_status()
    return r


def _column(groups, countries, group, country):
    """The index of the column under `group` (a merged header over several) whose country header is `country`."""
    current = None
    for j, (g, c) in enumerate(zip(groups, countries)):
        current = g.strip() if isinstance(g, str) else current
        if current == group and isinstance(c, str) and c.strip() == country:
            return j
    raise RuntimeError(f"no {country!r} column under {group!r}")


def parse_vintage(ws, group, country):
    """(the vintage's last quarter start, its value there) from one vintage sheet."""
    rows = [r for r in ws.iter_rows(values_only=True) if r]
    header = next((i for i, r in enumerate(rows) if isinstance(r[0], str) and r[0].strip() == "Date"), None)
    if header is None or header == 0:
        raise RuntimeError(f"sheet {ws.title}: no Date header row under a row of groups")
    col = _column(rows[header - 1], rows[header], group, country)
    dated = [r for r in rows[header + 1:] if isinstance(r[0], datetime.datetime)]
    quarter = pd.Period(ws.title, freq="Q")
    last = dated[-1] if dated else None
    if last is None or pd.Timestamp(last[0]) != quarter.start_time:
        raise RuntimeError(f"sheet {ws.title}: the last dated row is {last and last[0]}, not {quarter}")
    if not isinstance(last[col], (int, float)):
        raise RuntimeError(f"sheet {ws.title}: {country} {group} in {quarter} is {last[col]!r}")
    return quarter.start_time, float(last[col])


def parse_vintages(data, group, country):
    """Every vintage sheet ("YYYYQn") of a real-time workbook: date (the last quarter's start), value."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = [n for n in wb.sheetnames if re.fullmatch(r"\d{4}Q[1-4]", n)]
    if not names:
        raise RuntimeError("no vintage sheets (YYYYQn) in the workbook")
    df = pd.DataFrame([parse_vintage(wb[n], group, country) for n in names], columns=["date", "value"])
    return df.sort_values("date").reset_index(drop=True)


def published(dates, lag_days):
    """Each vintage's release: its quarter's last day plus `lag_days`, or the `RELEASED_LATE` day if later."""
    quarters = pd.PeriodIndex(pd.DatetimeIndex(dates), freq="Q")
    rule = quarters.end_time.normalize() + pd.Timedelta(days=lag_days)
    late = pd.to_datetime([RELEASED_LATE.get(str(q)) for q in quarters])
    return rule.where(~(late > rule), late)


class Hlw(Source):
    """One real-time HLW series (`SERIES`): one row per vintage, keyed by its last quarter's start."""

    name = "nyfed"

    def __init__(self, lag_days=65):
        self.lag_days = lag_days
        self._parsed = {}

    def fetch(self, series, start, end):
        if series not in self._parsed:
            workbook, group, country = SERIES[series]
            self._parsed[series] = parse_vintages(_get(DATA + workbook).content, group, country)
        df = self._parsed[series]
        df = df[(df["date"] >= pd.Timestamp(start)) & (df["date"] <= pd.Timestamp(end))].reset_index(drop=True)
        return df.assign(published=published(df["date"], self.lag_days))
