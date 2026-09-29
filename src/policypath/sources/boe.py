"""Bank of England data: daily series from the Interactive Statistical Database, and the fitted OIS and gilt curves.

Two sources, both under the cache namespace ``boe``:

- `BoeSeries`: one IADB series a day (SONIA ``IUDSOIA``, Bank Rate ``IUDBEDR``)
  over the IADB's CSV export. ``published`` is `lag_bdays` London business
  days after the date: SONIA for day d is published at 9:00 the next business
  day; the Bank Rate in force on d was announced by noon on d at the latest.
- `BoeCurve`: the Bank's daily fitted spot curves, continuously compounded,
  keyed by (date, tenor in months). ``OIS_SPOT`` is the OIS curve's short end,
  monthly from 1 to 60 months: the market path. ``GLC_SPOT`` is the nominal gilt
  curve's long end, half-yearly from 0.5 to 25 years (to 40 from 2016): the gilt
  legs. Published as Excel workbooks inside zip archives (one per span of years,
  plus the current month), and aimed to be out by noon the next business day, so
  ``published`` = the next London business day.

The workbooks change layout over the archive: OIS 2009-15 has no separate
short-end sheet (its "spot curve" is the monthly one), later ones do, and each
carries banner, formula and "Refresh" rows above the data. Sheets are picked by
"spot" in the name and their first maturity (1/12 or 0.5 years), not by their
full name, and only rows dated in column A are data. Blank cells are dropped:
holiday rows are blank throughout, and the GLC long end has no 0.5y rate on
20-50% of days in each year 2013-25 (none in 2009-12), so a day's curve starts
at 0.5y or 1y. Its 2y and 10y are there on every day it has a curve.

Both curves share ``latest-yield-curve-data.zip``: one `BoeCurve` downloads each
archive once, and parses a workbook only if its years reach the start asked for
(the GLC archive goes back to 1979: 1979-2004 are five workbooks of 4-5MB each).
"""

import datetime
import io
import re
import zipfile
import numpy as np
import openpyxl
import pandas as pd
import requests
from policypath.calendars import UK_BDAY
from policypath.sources.base import Source

IADB = "https://www.bankofengland.co.uk/boeapps/database/_iadb-fromshowcolumns.asp"
CURVES = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/"
LATEST = "latest-yield-curve-data.zip"      # every curve's current-month workbook
# Per curve: the archives holding it, the prefix its workbooks' names share, and its spot sheet.
CURVE_FILES = {"OIS_SPOT": (["oisddata.zip", LATEST], "OIS daily data", "monthly"),
               "GLC_SPOT": (["glcnominalddata.zip", LATEST], "GLC Nominal daily data", "half-yearly")}
# A spot sheet's first maturity in years: the short end is monthly, the long end half-yearly.
FIRST_MATURITY = {"monthly": 1 / 12, "half-yearly": 0.5}
HEADERS = {"User-Agent": "Mozilla/5.0 (policypath research)"}


def _get(url, **params):
    r = requests.get(url, params=params, headers=HEADERS, timeout=120)
    r.raise_for_status()
    return r


def published_after(dates, lag_bdays):
    """`lag_bdays` London business days after each date, never before the date itself."""
    days = pd.DatetimeIndex(dates).to_numpy().astype("datetime64[D]")
    lagged = np.busday_offset(days, lag_bdays, roll="backward", busdaycal=UK_BDAY.calendar)
    return pd.to_datetime(np.maximum(lagged, days))


def iadb(series, start, end):
    """One IADB series as date, value; days the database leaves blank are dropped."""
    r = _get(IADB, **{"csv.x": "yes", "Datefrom": pd.Timestamp(start).strftime("%d/%b/%Y"),
                      "Dateto": pd.Timestamp(end).strftime("%d/%b/%Y"), "SeriesCodes": series,
                      "CSVF": "TN", "UsingCodes": "Y", "VPD": "Y", "VFD": "N"})
    df = pd.read_csv(io.StringIO(r.text))
    if list(df.columns) != ["DATE", series]:
        raise RuntimeError(f"IADB {series}: unexpected columns {list(df.columns)}")
    df = df.dropna(subset=[series])
    return pd.DataFrame({"date": pd.to_datetime(df["DATE"], format="%d %b %Y"),
                         "value": df[series].astype(float)}).reset_index(drop=True)


class BoeSeries(Source):
    """A daily IADB series, keyed by date."""

    name = "boe"
    refetch_days = 7

    def __init__(self, lag_bdays=1):
        self.lag_bdays = lag_bdays

    def fetch(self, series, start, end):
        df = iadb(series, start, end)
        return df.assign(published=published_after(df["date"], self.lag_bdays))


def _maturities(ws, first):
    """The row of maturities in years, if this sheet's first maturity is `first` years."""
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        if row and isinstance(row[0], str) and row[0].strip().lower().startswith("years"):
            years = [v for v in row[1:] if isinstance(v, (int, float))]
            if years and abs(years[0] - first) < 1e-6:
                return years
    return None


def parse_curve_workbook(data, sheet="monthly"):
    """One spot sheet of a Bank of England curve workbook, long: date, tenor (months), value.

    `sheet` is a `FIRST_MATURITY` key: the monthly short end or the half-yearly long end.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for name in wb.sheetnames:
        if "spot" not in name.lower():
            continue
        ws = wb[name]
        years = _maturities(ws, FIRST_MATURITY[sheet])
        if years is None:          # the other end of the curve
            continue
        tenors = [round(y * 12) for y in years]
        rows = [r for r in ws.iter_rows(min_row=1, values_only=True)
                if r and isinstance(r[0], datetime.datetime)]
        wide = pd.DataFrame([r[1:len(tenors) + 1] for r in rows], columns=tenors,
                            index=pd.DatetimeIndex([r[0] for r in rows], name="date"))
        long = wide.apply(pd.to_numeric, errors="coerce").stack().rename("value").reset_index()
        long = long.rename(columns={"level_1": "tenor"}).dropna(subset=["value"])  # holidays: blank rows
        return long.astype({"tenor": int})
    raise RuntimeError(f"no {sheet} spot sheet in the workbook")


def last_year(member):
    """The last year a workbook covers, from its name ("..._2005 to 2015.xlsx"); None if it is open-ended."""
    m = re.search(r"_\d{4} to (\d{4})\.xlsx$", member)
    return int(m.group(1)) if m else None


class BoeCurve(Source):
    """One of the Bank's fitted spot curves (`CURVE_FILES`), keyed by date and tenor in months."""

    name = "boe"
    keys = ("date", "tenor")
    refetch_days = 7        # the current-month workbook can be refitted after first release

    def __init__(self, lag_bdays=1):
        self.lag_bdays = lag_bdays
        self._archives = {}     # archive name -> its zip, downloaded once
        self._sheets = {}       # (archive, workbook, sheet) -> the parsed spot sheet

    def _archive(self, archive):
        if archive not in self._archives:
            self._archives[archive] = zipfile.ZipFile(io.BytesIO(_get(CURVES + archive).content))
        return self._archives[archive]

    def _all(self, series, start):
        """Every workbook of the series whose years reach `start`, each parsed once; later files win a date.

        A workbook that ends before `start`'s year holds no date `fetch` keeps, so
        skipping it changes nothing but the time.
        """
        archives, prefix, sheet = CURVE_FILES[series]
        frames = []
        for archive in archives:
            z = self._archive(archive)
            for member in sorted(z.namelist()):
                if not (member.startswith(prefix) and member.endswith(".xlsx")):
                    continue
                last = last_year(member)
                if last is not None and last < start.year:
                    continue
                if (archive, member, sheet) not in self._sheets:
                    self._sheets[archive, member, sheet] = parse_curve_workbook(z.read(member), sheet)
                frames.append(self._sheets[archive, member, sheet])
        both = pd.concat(frames, ignore_index=True)
        return both.drop_duplicates(["date", "tenor"], keep="last")

    def fetch(self, series, start, end):
        start, end = pd.Timestamp(start), pd.Timestamp(end)
        df = self._all(series, start)
        df = df[(df["date"] >= start) & (df["date"] <= end)]
        df = df.sort_values(["date", "tenor"]).reset_index(drop=True)
        return df.assign(published=published_after(df["date"], self.lag_bdays))
